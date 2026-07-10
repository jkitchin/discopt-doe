"""Smoke tests for the ``discopt doe`` CLI surface.

Drives every verb through its pure ``do_*`` function (not via the
argparse layer) so the contract a future GUI binds to is what's
under test. Argparse parsing is exercised separately in one focused
test.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest

pytestmark = pytest.mark.smoke


openpyxl = pytest.importorskip("openpyxl")

from discopt.doe.cli import (  # noqa: E402
    DoEError,
    ExtendParams,
    NewParams,
    OptimizeParams,
    do_extend,
    do_fit,
    do_new,
    do_optimize,
    do_status,
    do_templates,
)
from discopt.doe.workbook import InputSpec, Workbook  # noqa: E402

# ──────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────


def _fill_response(path: Path, response: str, predictor) -> None:
    """Fill in the response column on every pending row using `predictor`."""
    wb = openpyxl.load_workbook(path)
    runs = wb["runs"]
    headers = [c.value for c in runs[1]]
    resp_idx = headers.index(response) + 1
    for row in runs.iter_rows(min_row=2):
        if row[0].value is None:
            continue
        row_dict = dict(zip(headers, [c.value for c in row]))
        row[resp_idx - 1].value = float(predictor(row_dict))
    wb.save(path)


# ──────────────────────────────────────────────────────────────────
# do_templates
# ──────────────────────────────────────────────────────────────────


def test_templates_lists_all():
    out = do_templates()
    names = [t["name"] for t in out["templates"]]
    assert names == [
        "linear",
        "polynomial-1d",
        "response-surface-2d",
        "response-surface-3d",
        "scheffe-linear",
        "scheffe-quadratic",
        "scheffe-special-cubic",
        "latin-square",
        "graeco-latin",
        "hyper-graeco-latin",
        "factorial-2level",
        "optimize",
    ]


# ──────────────────────────────────────────────────────────────────
# do_new — one test per template
# ──────────────────────────────────────────────────────────────────


def _new_params(tmpdir, *, template, inputs, n, degree=None, response="y", error=1.0, n_starts=1):
    return NewParams(
        output=Path(tmpdir) / f"{template}.xlsx",
        n=n,
        inputs=inputs,
        response_name=response,
        measurement_error=error,
        criterion="determinant",
        seed=0,
        n_starts=n_starts,
        template=template,
        degree=degree,
    )


def test_new_linear(tmp_path):
    out = do_new(_new_params(tmp_path, template="linear", inputs=[("x", 0.0, 10.0)], n=3))
    assert out["template"] == "linear"
    assert out["parameter_names"] == ["b0", "b1"]
    assert len(out["designs"]) == 3
    assert all(0 <= d["x"] <= 10 for d in out["designs"])
    assert out["next_command"].startswith("discopt doe status")


def test_new_polynomial_1d(tmp_path):
    out = do_new(
        _new_params(tmp_path, template="polynomial-1d", inputs=[("x", 0.0, 10.0)], n=4, degree=3)
    )
    assert out["parameter_names"] == ["b0", "b1", "b2", "b3"]
    assert len(out["designs"]) == 4


def test_new_response_surface_2d(tmp_path):
    out = do_new(
        _new_params(
            tmp_path,
            template="response-surface-2d",
            inputs=[("x1", 0.0, 10.0), ("x2", -5.0, 5.0)],
            n=6,
        )
    )
    assert out["parameter_names"] == ["b0", "b1", "b2", "b11", "b22", "b12"]
    assert len(out["designs"]) == 6


@pytest.mark.slow
def test_new_response_surface_3d(tmp_path):
    # Marked slow: a 10-parameter D-optimal search drives ~700 inner QP solves
    # and runs ~40s locally, blowing past the 120s CI fast-job timeout under
    # parallel load. n_starts=1 keeps the full-suite run as short as possible.
    # The response-surface model-based path is already covered by the 2D test.
    out = do_new(
        _new_params(
            tmp_path,
            template="response-surface-3d",
            inputs=[("x1", 0.0, 10.0), ("x2", -5.0, 5.0), ("x3", 0.0, 1.0)],
            n=10,
            n_starts=1,
        )
    )
    assert len(out["parameter_names"]) == 10
    assert len(out["designs"]) == 10


# ──────────────────────────────────────────────────────────────────
# Synthetic-truth round trip: new → fill → fit → recover
# ──────────────────────────────────────────────────────────────────


def test_fit_recovers_known_truth_response_surface(tmp_path):
    """do_fit recovers known quadratic-truth parameters from a filled workbook.

    The workbook is built directly with a fixed 3x3-grid-plus-center design
    rather than via do_new. This test validates *fit recovery*, not design
    optimality (the model-based design path is covered by
    test_new_response_surface_2d). do_new's inner D-optimal search costs ~20s
    even at n_starts=1 — that fixed cost is what previously pushed this test
    past the 120s CI timeout. Bypassing it keeps the run under a second.
    """
    inputs = [("x1", 0.0, 10.0), ("x2", -5.0, 5.0)]
    truth = {
        "b0": 1.0,
        "b1": 0.5,
        "b2": -0.3,
        "b11": -0.02,
        "b22": 0.1,
        "b12": 0.05,
    }

    def predict(row):
        x1, x2 = row["x1"], row["x2"]
        return (
            truth["b0"]
            + truth["b1"] * x1
            + truth["b2"] * x2
            + truth["b11"] * x1 * x1
            + truth["b22"] * x2 * x2
            + truth["b12"] * x1 * x2
        )

    # Fixed 3x3 grid (corners, edge midpoints, center) plus one extra center
    # run = 10 points: well-conditioned for the 6-term quadratic surface.
    design = [{"x1": x1, "x2": x2} for x1 in (0.0, 5.0, 10.0) for x2 in (-5.0, 0.0, 5.0)]
    design.append({"x1": 5.0, "x2": 0.0})

    wb_path = tmp_path / "response-surface-2d.xlsx"
    wb = Workbook.create(
        wb_path,
        template="response-surface-2d",
        template_args={},
        input_specs=[InputSpec(name, lb, ub) for name, lb, ub in inputs],
        criterion="determinant",
        measurement_error=0.01,
        seed=0,
        response_name="y",
    )
    wb.append_runs(1, design)
    wb.save()
    _fill_response(wb_path, "y", predict)

    fit = do_fit({"workbook": str(wb_path)})
    estimates = {p["name"]: p["estimate"] for p in fit["parameters"]}
    for name, expected in truth.items():
        assert estimates[name] == pytest.approx(expected, abs=1e-3)
    assert fit["n_observations"] == 10
    assert fit["log_det_fim"] > 0


def test_extend_appends_new_batch(tmp_path):
    inputs = [("x", 0.0, 10.0)]

    def predict(row):
        return 2.0 + 3.0 * row["x"]

    out = do_new(_new_params(tmp_path, template="linear", inputs=inputs, n=4, error=0.05))
    wb_path = Path(out["workbook_path"])
    _fill_response(wb_path, "y", predict)
    do_fit({"workbook": str(wb_path)})

    before_status = do_status({"workbook": str(wb_path)})
    assert before_status["n_pending"] == 0

    ext = do_extend(ExtendParams(workbook=wb_path, n=3, n_starts=1))
    assert ext["batch"] == 2
    assert len(ext["new_run_ids"]) == 3
    assert sorted(ext["new_run_ids"]) == ext["new_run_ids"]  # contiguous & sorted

    after_status = do_status({"workbook": str(wb_path)})
    assert after_status["n_completed"] == 4
    assert after_status["n_pending"] == 3
    assert after_status["n_total"] == 7


# ──────────────────────────────────────────────────────────────────
# Module escape hatch
# ──────────────────────────────────────────────────────────────────


_MODULE_EXPERIMENT_SOURCE = """
import discopt.modeling as dm
from discopt.estimate import Experiment, ExperimentModel


class CliTestKExp(Experiment):
    \"\"\"y = k * x for the --module escape-hatch test.\"\"\"

    def create_model(self, **kwargs):
        m = dm.Model("k_exp")
        k = m.continuous("k", lb=0.01, ub=20.0)
        x = m.continuous("x", lb=0.1, ub=10.0)
        return ExperimentModel(
            model=m,
            unknown_parameters={"k": k},
            design_inputs={"x": x},
            responses={"y": k * x},
            measurement_error={"y": 0.1},
        )
"""


@pytest.fixture
def module_experiment(tmp_path, monkeypatch):
    """Write a tiny Experiment module to a tmp dir and put it on sys.path."""
    import sys

    pkg_dir = tmp_path / "mod_pkg"
    pkg_dir.mkdir()
    (pkg_dir / "kexp.py").write_text(_MODULE_EXPERIMENT_SOURCE)
    monkeypatch.syspath_prepend(str(pkg_dir))
    sys.modules.pop("kexp", None)
    yield "kexp:CliTestKExp"
    sys.modules.pop("kexp", None)


def test_module_callable_escape_hatch(tmp_path, module_experiment):
    out_path = tmp_path / "mod.xlsx"
    out = do_new(
        NewParams(
            output=out_path,
            n=2,
            inputs=[("x", 0.1, 10.0)],
            response_name="y",
            measurement_error=0.1,
            criterion="determinant",
            seed=0,
            n_starts=3,
            module_callable=module_experiment,
            param_initial_guess={"k": 1.5},
        )
    )
    assert out["module_callable"] == module_experiment
    assert out["template"] is None
    assert out["parameter_names"] == ["k"]
    assert out["designs"][0]["x"] == pytest.approx(10.0, abs=1e-3)


def test_module_fit_refused(tmp_path, module_experiment):
    out_path = tmp_path / "mod.xlsx"
    do_new(
        NewParams(
            output=out_path,
            n=2,
            inputs=[("x", 0.1, 10.0)],
            response_name="y",
            measurement_error=0.1,
            criterion="determinant",
            seed=0,
            n_starts=3,
            module_callable=module_experiment,
            param_initial_guess={"k": 1.5},
        )
    )
    # Fill responses so fit reaches the template-check guard, not the
    # no-data guard.
    _fill_response(out_path, "y", lambda row: 1.5 * row["x"])
    with pytest.raises(DoEError, match="not yet implemented"):
        do_fit({"workbook": str(out_path)})


# ──────────────────────────────────────────────────────────────────
# Failure modes
# ──────────────────────────────────────────────────────────────────


def test_new_rejects_wrong_input_count(tmp_path):
    with pytest.raises((DoEError, ValueError), match="exactly two inputs"):
        do_new(
            _new_params(
                tmp_path,
                template="response-surface-2d",
                inputs=[("x", 0.0, 10.0)],
                n=4,
            )
        )


def test_new_rejects_both_template_and_module(tmp_path):
    with pytest.raises(DoEError, match="either --template or --module"):
        do_new(
            NewParams(
                output=tmp_path / "x.xlsx",
                n=1,
                inputs=[("x", 0, 1)],
                response_name="y",
                measurement_error=1.0,
                criterion="determinant",
                seed=0,
                n_starts=2,
                template="linear",
                module_callable="some.mod:thing",
            )
        )


def test_fit_no_completed_runs(tmp_path):
    out = do_new(_new_params(tmp_path, template="linear", inputs=[("x", 0, 1)], n=2))
    with pytest.raises(DoEError, match="no completed runs"):
        do_fit({"workbook": out["workbook_path"]})


def test_status_missing_workbook(tmp_path):
    with pytest.raises(FileNotFoundError):
        do_status({"workbook": str(tmp_path / "nope.xlsx")})


# ──────────────────────────────────────────────────────────────────
# Workbook persistence round-trip
# ──────────────────────────────────────────────────────────────────


def test_workbook_metadata_roundtrip(tmp_path):
    # Metadata is written at workbook *creation*; this test asserts only the
    # persistence round-trip, not anything about the generated design. Build
    # the workbook directly rather than paying do_new's ~15s D-optimal search
    # (the model-based design path is covered by test_new_response_surface_2d).
    inputs = [("x1", 0.0, 10.0), ("x2", -5.0, 5.0)]
    wb_path = tmp_path / "response-surface-2d.xlsx"
    wb = Workbook.create(
        wb_path,
        template="response-surface-2d",
        template_args={},
        input_specs=[InputSpec(name, lb, ub) for name, lb, ub in inputs],
        criterion="determinant",
        measurement_error=0.5,
        seed=0,
        response_name="y",
    )
    wb.save()

    wb = Workbook.open(wb_path)
    assert wb.template_name() == "response-surface-2d"
    assert wb.response_name() == "y"
    assert wb.measurement_error() == 0.5
    specs = wb.input_specs()
    assert [(s.name, s.lb, s.ub) for s in specs] == [
        ("x1", 0.0, 10.0),
        ("x2", -5.0, 5.0),
    ]
    _, names = wb.rebuild_experiment()
    assert names == ["b0", "b1", "b2", "b11", "b22", "b12"]


def test_fim_persisted_and_used_by_extend(tmp_path):
    # Asserts the FIM is persisted by do_fit and positive-definite — not that
    # the design is D-optimal. Build a fixed, well-conditioned 3x3-grid design
    # directly (as in test_fit_recovers_known_truth_response_surface) to skip
    # do_new's ~18s D-optimal search, which this test does not exercise.
    inputs = [("x1", 0.0, 10.0), ("x2", -5.0, 5.0)]
    wb_path = tmp_path / "response-surface-2d.xlsx"
    wb = Workbook.create(
        wb_path,
        template="response-surface-2d",
        template_args={},
        input_specs=[InputSpec(name, lb, ub) for name, lb, ub in inputs],
        criterion="determinant",
        measurement_error=0.5,
        seed=0,
        response_name="y",
    )
    design = [{"x1": x1, "x2": x2} for x1 in (0.0, 5.0, 10.0) for x2 in (-5.0, 0.0, 5.0)]
    design.append({"x1": 5.0, "x2": 0.0})
    wb.append_runs(1, design)
    wb.save()

    def predict(row):
        return 1 + 0.5 * row["x1"] - 0.3 * row["x2"]

    _fill_response(wb_path, "y", predict)
    do_fit({"workbook": str(wb_path)})

    wb = Workbook.open(wb_path)
    fim_data = wb.read_fim()
    assert fim_data is not None
    fim, names = fim_data
    assert names == ["b0", "b1", "b2", "b11", "b22", "b12"]
    eigvals = np.linalg.eigvalsh(fim)
    assert (eigvals > 0).all()


# ──────────────────────────────────────────────────────────────────
# do_optimize — resolves settings from the workbook (issue #1)
# ──────────────────────────────────────────────────────────────────


def _new_optimize_workbook(tmp_path, *, direction):
    """Create an optimize-template workbook created with the given direction."""
    return do_new(
        NewParams(
            output=Path(tmp_path) / "opt.xlsx",
            n=3,
            inputs=[("x", -5.0, 5.0)],
            response_name="y",
            measurement_error=0.05,
            criterion="determinant",
            seed=0,
            n_starts=1,
            template="optimize",
            optimize_criterion=direction,
        )
    )


def test_do_optimize_uses_stored_direction(tmp_path):
    """A minimize workbook must be minimized when no --criterion is given.

    Regression for the P0 bug where do_optimize ignored the workbook's stored
    direction and always used the argparse default ('maximize').
    """
    pytest.importorskip("sklearn")
    wb_path = Path(tmp_path) / "opt.xlsx"
    _new_optimize_workbook(tmp_path, direction="minimize")
    _fill_response(wb_path, "y", lambda row: (row["x"] - 1.0) ** 2)

    out = do_optimize(OptimizeParams(workbook=wb_path))

    assert out["criterion"] == "minimize"
    assert out["warnings"] == []


def test_do_optimize_explicit_override_warns(tmp_path):
    """Passing --criterion that differs from the stored value warns but obeys."""
    pytest.importorskip("sklearn")
    wb_path = Path(tmp_path) / "opt.xlsx"
    _new_optimize_workbook(tmp_path, direction="minimize")
    _fill_response(wb_path, "y", lambda row: (row["x"] - 1.0) ** 2)

    out = do_optimize(OptimizeParams(workbook=wb_path, criterion="maximize"))

    assert out["criterion"] == "maximize"
    assert any("criterion overridden" in w for w in out["warnings"])


# ──────────────────────────────────────────────────────────────────
# Excel formula cells in the response column (issue #2)
# ──────────────────────────────────────────────────────────────────


def _formula_workbook(tmp_path):
    out = do_new(_new_params(tmp_path, template="linear", inputs=[("x", 0.0, 10.0)], n=3))
    return Path(out["workbook_path"])


def test_formula_response_without_cached_value_raises(tmp_path):
    """A response formula openpyxl can't resolve is a loud error, not a drop.

    Regression for the P0 bug where '=AVERAGE(...)' in a response cell was
    silently classified as a pending run and dropped from fit/anova/optimize.
    """
    wb_path = _formula_workbook(tmp_path)
    book = openpyxl.load_workbook(wb_path)
    runs = book["runs"]
    headers = [c.value for c in runs[1]]
    resp_col = headers.index("y") + 1
    for r in runs.iter_rows(min_row=2):
        if r[0].value is None:
            continue
        r[resp_col - 1].value = "=AVERAGE(1,2)"
        break
    book.save(wb_path)

    with pytest.raises(ValueError, match="formula"):
        Workbook.open(wb_path).completed_runs()


def test_formula_response_with_cached_value_is_read(tmp_path, monkeypatch):
    """A response formula with an Excel-cached value is read as that value."""
    wb_path = _formula_workbook(tmp_path)
    book = openpyxl.load_workbook(wb_path)
    runs = book["runs"]
    headers = [c.value for c in runs[1]]
    resp_col = headers.index("y") + 1
    data_rows = [r for r in runs.iter_rows(min_row=2) if r[0].value is not None]
    for r in data_rows[:-1]:
        r[resp_col - 1].value = 5.0
    data_rows[-1][resp_col - 1].value = "=1+2"  # Excel would cache 3.0
    book.save(wb_path)

    # Stub the data_only view as if Excel had computed the formula: read the
    # formula-preserving rows and substitute the cached result.
    def fake(self):
        sheet = self._wb["runs"]
        rows = [list(r) for r in sheet.iter_rows(min_row=2, values_only=True)]
        for row in rows:
            for j, v in enumerate(row):
                if isinstance(v, str) and v.startswith("="):
                    row[j] = 3.0
        return rows

    monkeypatch.setattr(Workbook, "_cached_runs_rows", fake)

    completed = Workbook.open(wb_path).completed_runs()
    assert sorted(float(r["y"]) for r in completed) == [3.0, 5.0, 5.0]


# ──────────────────────────────────────────────────────────────────
# In-place save safety: backup + chart/image warning (issue #10)
# ──────────────────────────────────────────────────────────────────


def test_save_writes_one_time_backup(tmp_path):
    out = do_new(_new_params(tmp_path, template="linear", inputs=[("x", 0.0, 10.0)], n=3))
    wb_path = Path(out["workbook_path"])
    original = wb_path.read_bytes()

    wb = Workbook.open(wb_path)
    wb.append_runs(2, [{"x": 5.0}])
    wb.save()

    bak = wb_path.with_name(wb_path.name + ".bak")
    assert bak.exists()
    assert bak.read_bytes() == original


def test_open_warns_on_embedded_chart(tmp_path):
    from openpyxl.chart import BarChart, Reference

    out = do_new(_new_params(tmp_path, template="linear", inputs=[("x", 0.0, 10.0)], n=3))
    wb_path = Path(out["workbook_path"])
    book = openpyxl.load_workbook(wb_path)
    chart = BarChart()
    chart.add_data(Reference(book["runs"], min_col=1, min_row=1, max_row=2))
    book["runs"].add_chart(chart, "H2")
    book.save(wb_path)

    with pytest.warns(UserWarning, match="charts or images"):
        Workbook.open(wb_path)


# ──────────────────────────────────────────────────────────────────
# Falsy metadata round-trips; --error validation (issue #11)
# ──────────────────────────────────────────────────────────────────


def test_seed_zero_survives_round_trip(tmp_path):
    out = do_new(_new_params(tmp_path, template="linear", inputs=[("x", 0.0, 10.0)], n=3))
    wb_path = Path(out["workbook_path"])
    wb = Workbook.open(wb_path)
    # _new_params uses seed=0; it must not be silently replaced with 42.
    assert wb.seed() == 0


def test_measurement_error_half_survives(tmp_path):
    out = do_new(
        _new_params(tmp_path, template="linear", inputs=[("x", 0.0, 10.0)], n=3, error=0.5)
    )
    wb = Workbook.open(Path(out["workbook_path"]))
    assert wb.measurement_error() == 0.5


def test_cli_new_rejects_nonpositive_error(tmp_path, capsys):
    import argparse

    from discopt.doe.cli import add_subparser

    top = argparse.ArgumentParser()
    add_subparser(top.add_subparsers(dest="cmd"))
    args = top.parse_args(
        [
            "doe",
            "new",
            "linear",
            "-o",
            str(tmp_path / "z.xlsx"),
            "--input",
            "x:0:1",
            "--error",
            "0",
        ]
    )
    assert args.doe_func(args) == 1
    assert not (tmp_path / "z.xlsx").exists()


# ──────────────────────────────────────────────────────────────────
# --json output is always valid JSON (issue #36)
# ──────────────────────────────────────────────────────────────────


def test_json_new_factorial_is_valid_json(tmp_path, capsys):
    """`new factorial-2level --json` embeds a NaN criterion_value; the emitted
    JSON must still parse strictly (no bare NaN)."""
    import argparse
    import json

    from discopt.doe.cli import add_subparser

    top = argparse.ArgumentParser()
    add_subparser(top.add_subparsers(dest="cmd"))
    args = top.parse_args(
        [
            "doe",
            "new",
            "factorial-2level",
            "-o",
            str(tmp_path / "f.xlsx"),
            "--factor",
            "A:-1:1",
            "--factor",
            "B:-1:1",
            "--json",
        ]
    )
    assert args.doe_func(args) == 0
    out = capsys.readouterr().out
    # Strict parse: parse_constant fires on NaN/Infinity, so raise if present.
    def _boom(x):
        raise ValueError(f"non-finite literal {x!r} in JSON")

    parsed = json.loads(out, parse_constant=_boom)
    assert parsed["template"] == "factorial-2level"

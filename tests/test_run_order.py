"""The runs sheet's ``run_order`` column, and user-supplied bookkeeping columns.

A design is randomized so that drift cannot bias the factor estimates, and the
order is stored so that drift can be *looked for*. That only holds while the
stored order is the order that actually happened: a bench that runs #7 before
#3 and has nowhere to say so silently destroys the second half of the bargain.
"""

from __future__ import annotations

import openpyxl
import pytest
from discopt.doe.cli import DoEError, NewParams, do_new
from discopt.doe.workbook import Workbook


def _campaign(tmp_path, **kw) -> Workbook:
    params = dict(
        output=tmp_path / "w.xlsx",
        n=4,
        inputs=[("x", 0.0, 10.0)],
        response_name="y",
        measurement_error=1.0,
        criterion="determinant",
        seed=0,
        n_starts=2,
        template="linear",
    )
    params.update(kw)
    do_new(NewParams(**params))
    return Workbook.open(tmp_path / "w.xlsx")


def _set(path, run_id: int, column: str, value) -> None:
    """Write one cell by header name, as a person would in a spreadsheet."""
    book = openpyxl.load_workbook(path)
    sheet = book["runs"]
    header = [c.value for c in sheet[1]]
    col = header.index(column) + 1
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == run_id:
            sheet.cell(row=row[0].row, column=col, value=value)
    book.save(path)


def test_a_new_workbook_has_the_column_and_defaults_to_run_id(tmp_path) -> None:
    wb = _campaign(tmp_path)
    assert "run_order" in wb._runs_headers()
    # Blank everywhere: the runs went in the order they are written in.
    assert wb.run_order() == {1: 1, 2: 2, 3: 3, 4: 4}


def test_the_bench_can_record_a_deviation(tmp_path) -> None:
    path = tmp_path / "w.xlsx"
    _campaign(tmp_path)
    _set(path, 3, "run_order", 1)
    _set(path, 1, "run_order", 3)

    order = Workbook.open(path).run_order()
    assert order[3] == 1 and order[1] == 3
    assert order[2] == 2 and order[4] == 4  # untouched rows keep their place


def test_a_partly_filled_column_still_reads(tmp_path) -> None:
    """Only the runs that moved need a number."""
    path = tmp_path / "w.xlsx"
    _campaign(tmp_path)
    _set(path, 4, "run_order", 1)
    order = Workbook.open(path).run_order()
    assert order[4] == 1
    assert order[2] == 2


def test_two_runs_claiming_the_same_position_is_refused(tmp_path) -> None:
    path = tmp_path / "w.xlsx"
    _campaign(tmp_path)
    _set(path, 1, "run_order", 2)
    _set(path, 3, "run_order", 2)
    with pytest.raises(ValueError, match="both claim run_order 2"):
        Workbook.open(path).run_order()


def test_a_non_numeric_position_says_which_run(tmp_path) -> None:
    path = tmp_path / "w.xlsx"
    _campaign(tmp_path)
    _set(path, 2, "run_order", "second")
    with pytest.raises(ValueError, match="run 2 has run_order"):
        Workbook.open(path).run_order()


def test_run_order_appears_on_every_run_dict(tmp_path) -> None:
    """It is an ordinary column, so it travels with the rows."""
    path = tmp_path / "w.xlsx"
    _campaign(tmp_path)
    _set(path, 2, "run_order", 1)
    rows = {int(r["run_id"]): r for r in Workbook.open(path).all_runs()}
    assert rows[2]["run_order"] == 1
    assert rows[3]["run_order"] is None


def test_a_workbook_written_before_the_column_existed_still_reads(tmp_path) -> None:
    """Older files have no such column; the order is then the run_id order."""
    path = tmp_path / "w.xlsx"
    _campaign(tmp_path)
    book = openpyxl.load_workbook(path)
    sheet = book["runs"]
    header = [c.value for c in sheet[1]]
    sheet.delete_cols(header.index("run_order") + 1)
    book.save(path)

    wb = Workbook.open(path)
    assert "run_order" not in wb._runs_headers()
    assert wb.run_order() == {1: 1, 2: 2, 3: 3, 4: 4}


# ---------------------------------------------------------------------------
# extra_columns
# ---------------------------------------------------------------------------


def test_extra_columns_reach_the_runs_sheet(tmp_path) -> None:
    wb = _campaign(tmp_path, extra_columns=["operator", "lot"])
    headers = wb._runs_headers()
    assert headers == ["run_id", "batch", "x", "operator", "lot", "run_order", "y", "measured_at"]
    # They sit between the factors and the response, where anova_report looks
    # for blocking factors, and they come back on every run.
    assert all("operator" in r for r in wb.all_runs())


def test_a_design_that_needs_replicate_keeps_it(tmp_path) -> None:
    """The factorial designs add "replicate" themselves; a user column joins it."""
    wb = _campaign(
        tmp_path,
        template="factorial-2level",
        inputs=[("x", 0.0, 10.0), ("z", 0.0, 4.0)],
        factor_pairs={"x": (0.0, 10.0), "z": (0.0, 4.0)},
        n=4,
        extra_columns=["operator"],
    )
    headers = wb._runs_headers()
    assert "replicate" in headers and "operator" in headers
    assert headers.index("replicate") < headers.index("run_order")


def test_a_column_that_clashes_with_a_built_in_is_refused(tmp_path) -> None:
    for name in ("run_id", "batch", "run_order", "measured_at", " "):
        with pytest.raises(DoEError, match="clash|blank"):
            _campaign(tmp_path, extra_columns=[name])
